"""Build the client deck: shaker and controller options for a project (draft .pptx).

Shaker data comes from config/shakers (same payload calculation as tools/shaker_comparison.py).
Prices, hours and the labour rate are assumptions in this script (search for RATE, a_items,
hours_b1, b1, b2): check them before the deck goes to a client.

usage: python tools/client_options_deck.py [out.pptx]   (default: shaker_controller_options_draft.pptx)
requires: pip install -e ".[report]"
"""

import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

import numpy as np

sys.path.insert(0, str(Path(__file__).parent))
from shaker_comparison import EXTRA, KEYS, payload_limit  # noqa: E402

from vcactl.config import load_shaker  # noqa: E402

OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("shaker_controller_options_draft.pptx")


def shaker_data() -> dict:
    """Ratings and the 0.01 g^2/Hz payload curve of each shaker, keyed like config/shakers."""
    out = {}
    for k in KEYS:
        s = load_shaker(k)
        half = EXTRA[k]["half_mm"] or s.displacement_pp_mm / 2
        f = np.geomspace(200, min(10000, s.f_max_hz), 25)
        out[k] = dict(name=EXTRA[k]["short"], f_min=s.f_min_hz, f_max=s.f_max_hz,
                      fs=s.force_sine_peak_n, fr=s.force_random_rms_n, a_s=s.accel_sine_peak_g,
                      a_r=s.accel_random_rms_g, d=s.displacement_pp_mm, v=s.velocity_peak_m_s,
                      m=s.moving_mass_kg, stat=s.payload_static_max_kg,
                      curve=[[float(x), payload_limit(s, half, 0.01, x)[0]] for x in f])
    return out


DATA = shaker_data()

# Palette: deep instrument navy dominates, signal orange as the single sharp accent,
# a cool teal as supporting tone.
NAVY = RGBColor(0x17, 0x24, 0x33)
INK = RGBColor(0x1E, 0x25, 0x2D)
MUTED = RGBColor(0x5A, 0x64, 0x70)
ORANGE = RGBColor(0xE0, 0x62, 0x1B)
TEAL = RGBColor(0x0E, 0x76, 0x7A)
CARD = RGBColor(0xEF, 0xF2, 0xF5)
LINE = RGBColor(0xC9, 0xD0, 0xD8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ICE = RGBColor(0xB9, 0xC7, 0xD6)
HEAD_FONT, BODY_FONT = "Cambria", "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = 13.333, 7.5
MX = 0.6                       # side margin
FIT_WARNINGS = []


# ----------------------------------------------------------------------------- helpers
def est_lines(text, width_in, size_pt):
    """Rough line count for Calibri/Cambria text (average glyph ~0.5 em)."""
    chars_per_line = max(1, int(width_in * 72 / (0.5 * size_pt)))
    lines = 0
    for para in text.split("\n"):
        lines += max(1, -(-len(para) // chars_per_line))
    return lines


def check_fit(name, text, w, h, size, spacing=1.2, pad=0.1):
    need = est_lines(text, w - 2 * pad, size) * size * spacing / 72 + 2 * pad
    if need > h + 0.02:
        FIT_WARNINGS.append(f"{name}: needs ~{need:.2f} in, box {h:.2f} in")


def bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def text(slide, x, y, w, h, runs, size=16, color=INK, font=BODY_FONT, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name="text", margin=0.05,
         spacing_after=0, italic=False):
    """runs: str, or list of paragraphs; each paragraph a str or list of (text, opts) tuples."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(margin))
    paras = runs if isinstance(runs, list) else [runs]
    plain = []
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(spacing_after)
        parts = para if isinstance(para, list) else [(para, {})]
        line = ""
        for t, opts in parts:
            r = p.add_run()
            r.text = t
            line += t
            f = r.font
            f.name = opts.get("font", font)
            f.size = Pt(opts.get("size", size))
            f.bold = opts.get("bold", bold)
            f.italic = opts.get("italic", italic)
            f.color.rgb = opts.get("color", color)
        plain.append(line)
    check_fit(name, "\n".join(plain), w, h, size, pad=margin)
    return tb


def box(slide, x, y, w, h, fill=CARD, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        name="box", radius=0.06):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.name = name
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    return s


def badge(slide, x, y, label, d=0.5, fill=ORANGE, size=16, name="badge"):
    """The deck's motif: a filled circle with a number or letter."""
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.name = name
    c.fill.solid()
    c.fill.fore_color.rgb = fill
    c.line.fill.background()
    c.shadow.inherit = False
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size, r.font.bold, r.font.name = Pt(size), True, BODY_FONT
    r.font.color.rgb = WHITE
    return c


def title(slide, t, sub=None):
    text(slide, MX, 0.45, W - 2 * MX, 0.75, t, size=32, font=HEAD_FONT, bold=True,
         color=NAVY, name="title")
    if sub:
        text(slide, MX, 1.18, W - 2 * MX, 0.45, sub, size=16, color=MUTED, name="subtitle")


def footer(slide, n, dark=False):
    col = ICE if dark else MUTED
    text(slide, MX, H - 0.45, 8, 0.3, "Draft for discussion  ·  indicative figures, excl. VAT",
         size=10, color=col, name="footer")
    text(slide, W - MX - 1, H - 0.45, 1, 0.3, str(n), size=10, color=col, align=PP_ALIGN.RIGHT,
         name="slide number")


def table(slide, x, y, w, rows, col_w, row_h=0.36, size=12, header_fill=NAVY,
          first_col_bold=True, bold_rows=(), name="table", shade=True):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w),
                                   Inches(row_h * len(rows)))
    shape.name = name
    tbl = shape.table
    # plain styling: no banding from the default style
    tblPr = tbl._tbl.tblPr
    tblPr.set("bandRow", "0")
    tblPr.set("firstRow", "0")
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = header_fill
            elif shade and i % 2 == 0:
                cell.fill.fore_color.rgb = CARD
            else:
                cell.fill.fore_color.rgb = WHITE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = val
            r.font.name = BODY_FONT
            r.font.size = Pt(size)
            r.font.color.rgb = WHITE if i == 0 else INK
            r.font.bold = i == 0 or (j == 0 and first_col_bold) or i in bold_rows
            check_fit(f"{name}[{i},{j}]", val, col_w[j], row_h, size, pad=0.06)
    return shape


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=2.0, name="arrow"):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    c.name = name
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return c


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


# ----------------------------------------------------------------------------- slides
n = 0


def new_slide(dark=False):
    global n
    n += 1
    s = prs.slides.add_slide(BLANK)
    bg(s, NAVY if dark else WHITE)
    return s


# 1. Title -------------------------------------------------------------------
s = new_slide(dark=True)
badge(s, MX, 1.55, "~", d=0.7, size=24, name="motif")
text(s, MX, 2.45, 11.5, 0.95, "Vibration testing for [project name]", size=44, font=HEAD_FONT,
     bold=True, color=WHITE, name="title")
text(s, MX, 3.5, 11.5, 0.6, "Shaker and controller options, with indicative costs",
     size=22, color=ICE, name="subtitle")
text(s, MX, 5.2, 11.5, 0.45, "[Client name]  ·  prepared by [our company]  ·  October 2026  ·  draft",
     size=14, color=ICE, name="meta")
notes(s, "Draft. Replace [project name], [client name] and [our company]. All prices are "
         "indicative and must be confirmed by quotes before this goes to the client.")

# 2. Agenda -----------------------------------------------------------------
s = new_slide()
title(s, "What this presentation covers")
steps = [("1", "Shakers", "Two B&K shakers we own, and two TIRA systems we could buy"),
         ("2", "Your requirements", "The test parameters that decide which shaker fits"),
         ("3", "Controller", "Buy a VR9700 controller, or use our own controller"),
         ("4", "Costs and next steps", "Indicative cost estimates for both routes")]
cw = (W - 2 * MX - 3 * 0.35) / 4
for i, (num, head, body) in enumerate(steps):
    x = MX + i * (cw + 0.35)
    box(s, x, 2.0, cw, 3.6, name=f"step {num}")
    badge(s, x + 0.3, 2.3, num, d=0.65, size=22, name=f"step {num} badge")
    text(s, x + 0.3, 3.15, cw - 0.6, 0.85, head, size=20, bold=True, color=NAVY,
         name=f"step {num} head")
    text(s, x + 0.3, 4.05, cw - 0.6, 1.4, body, size=15, color=INK, name=f"step {num} body")
footer(s, n)
notes(s, "The decision we need from you comes after step 2: your requirements decide the "
         "shaker, and the shaker plus your test types decide which controller route makes sense.")

# 3. What we have -------------------------------------------------------------
s = new_slide()
title(s, "What we already have", "Our starting point: two shakers, a DAQ device and our own controller software")
cards = [
    ("B&K 4809 vibration exciter", "Small exciter: 44.5 N, 10 Hz to 20 kHz. We have no amplifier "
     "for it; a B&K 2718 power amplifier would have to be bought.", "Needs amplifier", ORANGE),
    ("B&K 4801/4812 with 2707 amplifier", "Complete system with blower: 445 N, 5 Hz to 10 kHz. "
     "Needs three-phase mains (380 V) at the test location.", "Ready to use", TEAL),
    ("NI USB-6211 DAQ device", "Generates the drive signal and reads the accelerometer. It "
     "cannot power an IEPE accelerometer: that needs a small signal conditioner.",
     "Needs IEPE conditioner", ORANGE),
    ("Frame-based controller software", "Our own closed-loop random vibration controller. "
     "Runs and is tested in simulation; commissioning on the real hardware is still to do.",
     "Needs commissioning", ORANGE),
]
cw, ch = (W - 2 * MX - 0.4) / 2, 2.25
for i, (head, body, chip, chip_col) in enumerate(cards):
    x = MX + (i % 2) * (cw + 0.4)
    y = 1.95 + (i // 2) * (ch + 0.35)
    box(s, x, y, cw, ch, name=f"card {i + 1}")
    badge(s, x + 0.3, y + 0.3, str(i + 1), d=0.5, size=16, name=f"card {i + 1} badge")
    text(s, x + 1.0, y + 0.28, cw - 1.3, 0.5, head, size=18, bold=True, color=NAVY,
         name=f"card {i + 1} head")
    text(s, x + 1.0, y + 0.82, cw - 1.3, 0.95, body, size=14, name=f"card {i + 1} body")
    chip_w = 2.6
    c = box(s, x + 1.0, y + ch - 0.5, chip_w, 0.34, fill=None, line=chip_col,
            name=f"card {i + 1} chip", radius=0.5)
    text(s, x + 1.0, y + ch - 0.5, chip_w, 0.34, chip, size=12, bold=True, color=chip_col,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, name=f"card {i + 1} chip text",
         margin=0.02)
footer(s, n)

# 4. Shaker comparison table ---------------------------------------------------
s = new_slide()
title(s, "Four shakers compared", "Two we own, two complete TIRA systems we could buy")
d = DATA
order = ["bk4809", "bk4801_4812", "tv51110", "tv52110"]


def fmt_f(f):
    return f"{f / 1000:g} kHz" if f >= 1000 else f"{f:g} Hz"


rows = [["", "B&K 4809", "B&K 4801/4812", "TIRA TV 51110", "TIRA TV 52110"],
        ["Availability", "Owned, no amplifier", "Owned, complete", "To buy", "To buy"],
        ["Frequency range", *[f"{fmt_f(d[k]['f_min'])} – {fmt_f(d[k]['f_max'])}" for k in order]],
        ["Force, sine peak", *[f"{d[k]['fs']:g} N" for k in order]],
        ["Force, random rms", *[f"{d[k]['fr']:g} N" + (" *" if k.startswith("bk") else "")
                                for k in order]],
        ["Acceleration, sine / random", *[f"{d[k]['a_s']:g} g / {d[k]['a_r']:g} g rms" for k in order]],
        ["Stroke, peak-peak", *[f"{d[k]['d']:g} mm" for k in order]],
        ["Velocity, peak", *[f"{d[k]['v']:g} m/s" for k in order]],
        ["Moving mass", *[f"{d[k]['m'] * 1000:g} g" for k in order]],
        ["Static payload limit, vertical", *[f"{d[k]['stat']:g} kg" for k in order]],
        ["Still to buy", "B&K 2718 amplifier", "Nothing (3-phase mains)", "Complete system",
         "Complete system"]]
cols = [3.35, 2.25, 2.25, 2.14, 2.14]
table(s, MX, 1.85, sum(cols), rows, cols, row_h=0.4, size=13, name="shaker table",
      bold_rows=())
text(s, MX, 6.3, W - 2 * MX, 0.55, "* Derived from the sine rating (the B&K manuals give no "
     "random rating). Static payload limit: the weight that uses up half the stroke on a "
     "vertically mounted shaker.", size=11, color=MUTED, name="table note")
footer(s, n)
notes(s, "Values from the manufacturers' datasheets and manuals, as configured in our "
         "controller's shaker files. Random ratings for the B&K shakers are sine peak / sqrt(2).")

# 5. Payload chart ------------------------------------------------------------
s = new_slide()
title(s, "How much payload each shaker can carry",
      "The wider the frequency band, the more force a test needs, and the less payload is left")
cd = XyChartData()
series_style = {"bk4801_4812": (ORANGE, XL_MARKER_STYLE.DIAMOND),
                "tv51110": (NAVY, XL_MARKER_STYLE.CIRCLE),
                "tv52110": (RGBColor(0x7A, 0x86, 0x94), XL_MARKER_STYLE.SQUARE),
                "bk4809": (TEAL, XL_MARKER_STYLE.TRIANGLE)}
for k in ["bk4801_4812", "tv51110", "tv52110", "bk4809"]:
    ser = cd.add_series(d[k]["name"])
    for f, m in d[k]["curve"]:
        ser.add_data_point(f, round(m, 3))
gf = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, Inches(MX), Inches(1.85),
                        Inches(8.1), Inches(4.75), cd)
gf.name = "payload chart"
ch = gf.chart
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.BOTTOM
ch.legend.include_in_layout = False
ch.legend.font.size = Pt(12)
ch.legend.font.color.rgb = INK
ch.font.name = BODY_FONT
ch.font.size = Pt(12)
ch.font.color.rgb = MUTED
for plot in ch.plots:
    for ser in plot.series:
        key = next(k for k in d if d[k]["name"] == ser.name)
        col, mk = series_style[key]
        ser.format.line.color.rgb = col
        ser.format.line.width = Pt(2.25)
        ser.smooth = False
        ser.marker.style = XL_MARKER_STYLE.NONE
xa, ya = ch.category_axis, ch.value_axis
for ax, lo, hi, label in ((xa, 200, 10000, "Highest test frequency (Hz)"),
                          (ya, 0.2, 20, "Max payload (kg)")):
    ax.minimum_scale, ax.maximum_scale = lo, hi
    scaling = ax._element.find(qn("c:scaling"))
    lb = etree.SubElement(scaling, qn("c:logBase"))
    lb.set("val", "10")
    scaling.insert(0, lb)
    ax.has_major_gridlines = True
    ax.major_gridlines.format.line.color.rgb = RGBColor(0xE2, 0xE6, 0xEA)
    ax.format.line.color.rgb = LINE
    ax.has_title = True
    ax.axis_title.text_frame.text = label
    tr = ax.axis_title.text_frame.paragraphs[0].runs[0]
    tr.font.size, tr.font.bold, tr.font.color.rgb = Pt(12), False, MUTED
    ax.tick_labels.font.size = Pt(12)
    ax.tick_labels.font.color.rgb = MUTED
xa.tick_labels.number_format = "#,##0"
xa.tick_labels.number_format_is_linked = False
ya.tick_labels.number_format = "0.0#"
ya.tick_labels.number_format_is_linked = False

rx, rw = 9.15, W - MX - 9.15
callouts = [("12.6 kg", "B&K 4801/4812 up to 500 Hz, still 2.8 kg at 10 kHz", ORANGE),
            ("0.4 – 4.9 kg", "TIRA systems, 200 Hz to 7 kHz", NAVY),
            ("0.3 – 2.3 kg", "B&K 4809: light specimens only, but up to 20 kHz", TEAL)]
for i, (big, small, col) in enumerate(callouts):
    y = 1.95 + i * 1.45
    text(s, rx, y, rw, 0.6, big, size=30, bold=True, color=col, font=HEAD_FONT,
         name=f"callout {i + 1} value")
    text(s, rx, y + 0.62, rw, 0.7, small, size=14, color=INK, name=f"callout {i + 1} label")
text(s, MX, 6.62, W - 2 * MX, 0.4, "Payload = accelerometer + fixture + specimen. Flat random "
     "profile from 20 Hz at 0.01 g²/Hz (1.3 g rms at 200 Hz, 10 g rms at 10 kHz), vertical mounting.",
     size=11, color=MUTED, name="chart note")
footer(s, n)
notes(s, "Example profile only: your real profile can change these numbers a lot. A higher "
         "level or a lower start frequency reduces the payload. Our controller's pre-flight "
         "check computes the exact limit for any profile and payload.")

# 6. Which shaker for which test ------------------------------------------------
s = new_slide()
title(s, "Which shaker suits which test")
cols4 = [
    ("B&K 4809", TEAL, "Small, light specimens; high frequencies up to 20 kHz",
     "Low force (31.5 N rms random); 8 mm stroke", "B&K 2718 amplifier (quote needed)"),
    ("B&K 4801/4812", ORANGE, "Heavier specimens up to about 10 kg; high test levels",
     "80 kg exciter, three-phase mains, blower noise; older equipment", "None"),
    ("TIRA TV 51110", NAVY, "Compact complete system; more random force than the TV 52110",
     "Up to 7 kHz; payloads of a few kg at most", "Complete system (quote needed)"),
    ("TIRA TV 52110", RGBColor(0x5A, 0x66, 0x74), "More stroke (15 mm) and a stiffer suspension",
     "Least random force of the TIRAs (50 N rms); up to 7 kHz",
     "Complete system (quote needed)"),
]
cw = (W - 2 * MX - 3 * 0.3) / 4
for i, (name, col, best, watch, buy) in enumerate(cols4):
    x = MX + i * (cw + 0.3)
    box(s, x, 1.6, cw, 5.05, name=f"{name} card")
    text(s, x + 0.25, 1.85, cw - 0.5, 0.5, name, size=18, bold=True, color=col,
         name=f"{name} head")
    y = 2.5
    for label, body, hh in (("Best for", best, 1.15), ("Watch out", watch, 1.45),
                            ("Still to buy", buy, 0.95)):
        text(s, x + 0.25, y, cw - 0.5, 0.35, label, size=12, bold=True, color=col,
             name=f"{name} {label}")
        text(s, x + 0.25, y + 0.34, cw - 0.5, hh - 0.35, body, size=14,
             name=f"{name} {label} text")
        y += hh + 0.1
footer(s, n)

# 7. Questions for the client ---------------------------------------------------
s = new_slide()
title(s, "What we need to know from you", "Your answers decide the shaker and the controller route")
qs = [("Payload", "Total mass of specimen and fixture? Size and mounting points?"),
      ("Frequency range", "Lowest and highest test frequency?"),
      ("Vibration level", "Test profile (PSD, g rms), or the standard it comes from?"),
      ("Stroke", "Low-frequency content or displacement limits in the test?"),
      ("Test types", "Random only, or also sine sweeps, shock or sine-on-random?"),
      ("Axes and orientation", "Vertical only, or horizontal axes too?"),
      ("Control points", "One control accelerometer, or several (averaging, limiting)?"),
      ("Duration and volume", "Test length, number of specimens, how often?"),
      ("Documentation", "Test reports, traceable calibration, accreditation (ISO/IEC 17025)?"),
      ("Site, timeline, budget", "Mains (three-phase?), space, noise; when, and what budget?")]
cw = (W - 2 * MX - 0.4) / 2
for i, (head, q) in enumerate(qs):
    col, row = i // 5, i % 5
    x = MX + col * (cw + 0.4)
    y = 1.75 + row * 1.0
    box(s, x, y, cw, 0.9, name=f"question {i + 1}")
    badge(s, x + 0.18, y + 0.2, str(i + 1), d=0.5, size=15, name=f"question {i + 1} badge")
    text(s, x + 0.85, y + 0.04, cw - 1.0, 0.32, head, size=14, bold=True, color=NAVY,
         name=f"question {i + 1} head")
    text(s, x + 0.85, y + 0.33, cw - 1.0, 0.56, q, size=14, color=INK,
         name=f"question {i + 1} text")
footer(s, n)
notes(s, "Questions 1-4 decide the shaker. Questions 5, 7 and 9 decide the controller route: "
         "our own controller does random tests with one control accelerometer today.")

# 8. Two controller routes --------------------------------------------------------
s = new_slide()
title(s, "Two routes to a working test system", "Either way: same shaker, IEPE accelerometer and specimen fixture")
routes = [("A", "Buy a controller: VR9700", ORANGE,
           "Commercial vibration controller from Vibration Research. We integrate it with "
           "the chosen shaker and train your staff.",
           ["Proven product, ready to use", "Several test types (licensed software modules)",
            "Vendor support and built-in reports"],
           ["About €25,000 for the controller", "Features fixed by the vendor"]),
          ("B", "Use our own controller", TEAL,
           "Our frame-based controller software on a PC with the NI USB-6211. We commission, "
           "validate and support it.",
           ["Lower cost; mostly hardware we own", "Fully adaptable to your project",
            "Safety checks built around your shaker"],
           ["Random tests with one control point today", "Commissioning and validation first"])]
cw = (W - 2 * MX - 0.4) / 2
for i, (letter, head, col, desc, pros, cons) in enumerate(routes):
    x = MX + i * (cw + 0.4)
    box(s, x, 1.85, cw, 4.85, name=f"route {letter}")
    badge(s, x + 0.3, 2.1, letter, d=0.65, fill=col, size=22, name=f"route {letter} badge")
    text(s, x + 1.15, 2.15, cw - 1.4, 0.55, head, size=22, bold=True, color=NAVY,
         font=HEAD_FONT, name=f"route {letter} head")
    text(s, x + 0.3, 2.95, cw - 0.6, 0.9, desc, size=15, name=f"route {letter} desc")
    text(s, x + 0.3, 3.95, cw - 0.6, 0.35, "Strengths", size=13, bold=True, color=col,
         name=f"route {letter} strengths head")
    text(s, x + 0.3, 4.3, cw - 0.6, 1.1, [("+  " + p) for p in pros], size=15,
         name=f"route {letter} strengths", spacing_after=3)
    text(s, x + 0.3, 5.45, cw - 0.6, 0.35, "Limitations", size=13, bold=True, color=col,
         name=f"route {letter} limits head")
    text(s, x + 0.3, 5.8, cw - 0.6, 0.8, [("–  " + c) for c in cons], size=15,
         name=f"route {letter} limits", spacing_after=3)
footer(s, n)
notes(s, "VR9700 features depend on the licensed software modules: confirm with Vibration "
         "Research which modules the €25,000 estimate includes.")

# 9. Own controller: hardware -------------------------------------------------------
s = new_slide()
title(s, "Our own controller: the hardware", "A closed loop: drive the shaker, measure the response, correct the drive")
bw, bh = 2.35, 1.2
ytop, ybot = 2.25, 4.55
xs = [MX, MX + 3.15, MX + 6.3, MX + 9.45]
top = [("PC with controller software", NAVY, "Windows or Linux"),
       ("NI USB-6211 DAQ", NAVY, "owned"),
       ("Power amplifier", NAVY, "B&K 2707 (owned) or 2718"),
       ("Shaker", ORANGE, "plus fixture, specimen")]
for i, (head, col, sub) in enumerate(top):
    b = box(s, xs[i], ytop, bw, bh, fill=col, name=f"block {head}")
    text(s, xs[i] + 0.1, ytop + 0.08, bw - 0.2, 0.62, head, size=15, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, name=f"block {head} text")
    text(s, xs[i] + 0.1, ytop + 0.74, bw - 0.2, 0.36, sub, size=12, color=WHITE,
         align=PP_ALIGN.CENTER, name=f"block {head} sub")
labels = ["USB", "drive", "current"]
for i in range(3):
    arrow(s, xs[i] + bw + 0.05, ytop + bh / 2, xs[i + 1] - 0.05, ytop + bh / 2,
          name=f"arrow {i + 1}")
    text(s, xs[i] + bw, ytop - 0.32, xs[i + 1] - xs[i] - bw, 0.3, labels[i], size=11,
         color=MUTED, align=PP_ALIGN.CENTER, name=f"arrow {i + 1} label")
bottom = [("IEPE signal conditioner", TEAL, "to buy"), ("IEPE accelerometer", TEAL, "to buy; on the table")]
bx = [xs[2], xs[3]]
for i, (head, col, sub) in enumerate(bottom):
    box(s, bx[i], ybot, bw, bh, fill=col, name=f"block {head}")
    text(s, bx[i] + 0.1, ybot + 0.08, bw - 0.2, 0.62, head, size=15, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, name=f"block {head} text")
    text(s, bx[i] + 0.1, ybot + 0.74, bw - 0.2, 0.36, sub, size=12, color=WHITE,
         align=PP_ALIGN.CENTER, name=f"block {head} sub")
# shaker -> accelerometer (mounted on table), accelerometer -> conditioner, conditioner -> DAQ
arrow(s, xs[3] + bw / 2, ytop + bh + 0.05, xs[3] + bw / 2, ybot - 0.05, color=TEAL,
      name="arrow shaker to accelerometer")
arrow(s, xs[3] - 0.05, ybot + bh / 2, xs[2] + bw + 0.05, ybot + bh / 2, color=TEAL,
      name="arrow accelerometer to conditioner")
# conditioner -> DAQ: elbow drawn as two segments
c1 = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(xs[2] - 0.05), Inches(ybot + bh / 2),
                            Inches(xs[1] + bw / 2), Inches(ybot + bh / 2))
c1.line.color.rgb, c1.line.width, c1.name = TEAL, Pt(2), "line conditioner to DAQ"
arrow(s, xs[1] + bw / 2, ybot + bh / 2, xs[1] + bw / 2, ytop + bh + 0.05, color=TEAL,
      name="arrow to DAQ input")
text(s, xs[1] + bw / 2 + 0.1, ybot + bh / 2 - 0.36, 2.6, 0.3, "measured acceleration",
     size=11, color=TEAL, name="feedback label")
text(s, MX, 5.95, W - 2 * MX, 0.75,
     [[("Upgrade option: ", {"bold": True, "color": NAVY}),
       ("an NI USB-4431 replaces the USB-6211 and the conditioner: built-in IEPE supply, "
        "24-bit resolution and anti-aliasing filters. Our software already supports it.", {})]],
     size=14, name="upgrade note")
footer(s, n)
notes(s, "The amplifier's own protections (current and displacement trips on the 2707, current "
         "limit on the 2718) stay active as a second safety layer next to the software limits.")

# 10. Own controller: software ------------------------------------------------------
s = new_slide()
title(s, "Our own controller: the software", "What our frame-based controller does today")
feats = [("Closed-loop random control", "Corrects the drive spectrum every 82 ms to follow "
          "the target profile (control band up to about 7.5 kHz)."),
         ("Pretest before every test", "Measures the noise floor and the shaker's response at a "
          "low level before any real test level."),
         ("Safe level ramp", "Starts 12 dB below the target and steps up 3 dB at a time, only "
          "when the previous level is under control."),
         ("Safety checks", "Before output: force, payload, stroke and velocity. During the "
          "test: overload, open loop and tolerance aborts."),
         ("Logging and plots", "Every run is logged: spectra against the tolerance bands, "
          "drive levels and status."),
         ("Simulator and configuration", "Tests can be prepared without hardware; shakers, DAQ "
          "devices and profiles are plain text files.")]
cw, ch = (W - 2 * MX - 2 * 0.35) / 3, 2.15
for i, (head, body) in enumerate(feats):
    x = MX + (i % 3) * (cw + 0.35)
    y = 1.9 + (i // 3) * (ch + 0.3)
    box(s, x, y, cw, ch, name=f"feature {i + 1}")
    badge(s, x + 0.25, y + 0.25, str(i + 1), d=0.45, fill=TEAL, size=14,
          name=f"feature {i + 1} badge")
    text(s, x + 0.85, y + 0.15, cw - 1.05, 0.68, head, size=16, bold=True, color=NAVY,
         anchor=MSO_ANCHOR.MIDDLE, name=f"feature {i + 1} head")
    text(s, x + 0.25, y + 0.9, cw - 0.5, 1.15, body, size=14, name=f"feature {i + 1} body")
footer(s, n)
notes(s, "Status: implemented and tested against a simulated shaker. Before first use it must be "
         "commissioned and validated on the real shaker. Not yet available: sine, shock, "
         "multi-point control; these can be developed if your tests need them.")

# 11. A vs B comparison -------------------------------------------------------------
s = new_slide()
title(s, "Bought or own controller: side by side")
rows = [["", "A: VR9700", "B: our own controller"],
        ["Test types", "Random, sine, shock and more (per licensed module)",
         "Random; other types developed on request"],
        ["Control points", "Several inputs; averaging and limit control",
         "One control accelerometer"],
        ["Readiness", "Commercial product, ready to use",
         "Ready in simulation; commissioning and validation needed"],
        ["Reporting", "Built-in test reports", "Run logs and plots; report template to develop"],
        ["Support", "Vendor support and updates", "Our team, with direct access to the developers"],
        ["Flexibility", "Vendor feature set", "Fully adaptable to your tests"],
        ["Indicative cost", "about €28,600", "about €13,400 – €17,800"]]
cols = [2.6, 4.97, 4.56]
table(s, MX, 1.6, sum(cols), rows, cols, row_h=0.6, size=14, name="route table",
      bold_rows=(7,))
footer(s, n)
notes(s, "VR9700 capabilities to be confirmed with Vibration Research for the modules in the "
         "quote. Costs: see the next three slides.")

# 12/13. Cost estimates -------------------------------------------------------------
RATE = 100


def eur(v):
    return f"€{v:,.0f}"


s = new_slide()
title(s, "Cost estimate A: VR9700", "Controller route only; shaker-related costs are on slide 14")
a_items = [("VR9700 controller with random control software", 25000, "estimate, quote needed"),
           ("IEPE accelerometer with cable", 1200, "estimate"),
           (f"Integration, commissioning and training: 24 h × €{RATE}/h", 24 * RATE, "our work")]
total_a = sum(v for _, v, _ in a_items)
rows = [["Item", "Cost", "Basis"]] + [[t, eur(v), b] for t, v, b in a_items] + \
       [["Total, route A", eur(total_a), "excl. VAT"]]
cols = [6.6, 1.8, 3.73]
table(s, MX, 1.85, sum(cols), rows, cols, row_h=0.55, size=15, name="cost table A",
      bold_rows=(len(rows) - 1,), first_col_bold=False)
box(s, MX, 5.1, sum(cols), 1.2, name="note A")
text(s, MX + 0.3, 5.2, sum(cols) - 0.6, 1.0,
     "Sine, shock or extra input channels may need additional VR9700 software modules or "
     "hardware. These are not included and depend on your answers to the questions on slide 7.",
     size=14, name="note A text", anchor=MSO_ANCHOR.MIDDLE)
footer(s, n)
notes(s, "Assumptions to check: VR9700 €25,000 (estimate given internally; confirm what it "
         f"includes), accelerometer €1,200, labour rate €{RATE}/h (placeholder: replace with our rate).")

s = new_slide()
title(s, "Cost estimate B: our own controller", "Two hardware variants; shaker-related costs are on slide 14")
hours_b1 = {"Commissioning on the shaker": 40, "Validation against a reference accelerometer": 24,
            "Test report template and documentation": 24, "Training": 8}
hours_b2 = dict(hours_b1, **{"Commissioning on the shaker": 32})
cont = 0.2
h1 = round(sum(hours_b1.values()) * (1 + cont))
h2 = round(sum(hours_b2.values()) * (1 + cont))
b1 = {"DAQ": 0, "IEPE": 700, "acc": 1200, "lab": h1 * RATE}
b2 = {"DAQ": 6000, "IEPE": 0, "acc": 1200, "lab": h2 * RATE}
total_b1, total_b2 = sum(b1.values()), sum(b2.values())
rows = [["Item", "B1: USB-6211 (owned)", "B2: USB-4431 upgrade"],
        ["DAQ device", "€0 (owned)", eur(b2["DAQ"]) + " (estimate)"],
        ["IEPE signal conditioner", eur(b1["IEPE"]) + " (estimate)", "€0 (built into DAQ)"],
        ["IEPE accelerometer with cable", eur(b1["acc"]) + " (estimate)", eur(b2["acc"]) + " (estimate)"],
        [f"Our work incl. 20 % contingency, at €{RATE}/h", f"{h1} h = {eur(b1['lab'])}",
         f"{h2} h = {eur(b2['lab'])}"],
        ["Total", eur(total_b1), eur(total_b2)]]
cols = [5.0, 3.56, 3.57]
table(s, MX, 1.85, sum(cols), rows, cols, row_h=0.55, size=15, name="cost table B",
      bold_rows=(len(rows) - 1,), first_col_bold=False)
text(s, MX, 5.3, sum(cols), 0.95,
     [[("Our work covers: ", {"bold": True, "color": NAVY}),
       ("commissioning on the shaker, validation against a reference accelerometer, a test "
        "report template, documentation and training. Not included: sine, shock or multi-point "
        "control (developed on request).", {})]],
     size=14, name="work note")
footer(s, n)
notes(s, f"Hours B1: {hours_b1}, +20% contingency = {h1} h. B2 saves commissioning time "
         f"(no external conditioner): {h2} h. Labour rate €{RATE}/h is a placeholder. "
         "USB-4431 €6,000 and IEPE conditioner €700 are indicative, to be confirmed by quotes. "
         "The existing controller software is not charged here: decide whether a licence fee "
         "applies.")

# 14. Overview -----------------------------------------------------------------------
s = new_slide()
title(s, "Cost overview", "Indicative totals for the controller route, plus shaker-dependent costs")
stats = [("A", "VR9700", total_a, ORANGE), ("B1", "Own controller, USB-6211", total_b1, TEAL),
         ("B2", "Own controller, USB-4431", total_b2, TEAL)]
cw = (W - 2 * MX - 2 * 0.35) / 3
for i, (lab, desc, val, col) in enumerate(stats):
    x = MX + i * (cw + 0.35)
    box(s, x, 1.85, cw, 2.15, name=f"stat {lab}")
    badge(s, x + 0.3, 2.1, lab, d=0.55, fill=col, size=15, name=f"stat {lab} badge")
    text(s, x + 1.0, 2.05, cw - 1.2, 0.68, desc, size=15, bold=True, color=NAVY,
         anchor=MSO_ANCHOR.MIDDLE, name=f"stat {lab} desc")
    text(s, x + 0.3, 2.9, cw - 0.6, 0.95, eur(val), size=44, bold=True, color=col,
         font=HEAD_FONT, name=f"stat {lab} value")
text(s, MX, 4.3, W - 2 * MX, 0.4, "Shaker-dependent costs, on top of either route",
     size=16, bold=True, color=NAVY, name="shaker costs head")
rows = [["Shaker", "Extra cost"],
        ["B&K 4801/4812 with 2707", "None; needs three-phase mains at the test location"],
        ["B&K 4809", "B&K 2718 power amplifier (quote needed)"],
        ["TIRA TV 51110 or TV 52110", "Complete system (quote needed)"],
        ["Any shaker", "Fixture for your specimen (designed once requirements are known)"]]
cols = [4.0, 8.13]
table(s, MX, 4.75, sum(cols), rows, cols, row_h=0.36, size=13, name="shaker cost table")
footer(s, n)

# 15. Next steps --------------------------------------------------------------------
s = new_slide(dark=True)
text(s, MX, 0.55, W - 2 * MX, 0.8, "Next steps", size=36, font=HEAD_FONT, bold=True,
     color=WHITE, name="title")
steps = [("1", "You", "Answer the questions on slide 7: payload, frequency range, levels, "
          "stroke, test types and documentation needs."),
         ("2", "We", "Select the shaker, request quotes (VR9700, accelerometer, amplifier or "
          "TIRA system) and firm up the estimates."),
         ("3", "Together", "Decide on the controller route; for route B, agree on a commissioning "
          "and validation plan.")]
for i, (num, who, what) in enumerate(steps):
    y = 1.9 + i * 1.5
    badge(s, MX, y, num, d=0.75, size=24, name=f"next {num} badge")
    text(s, MX + 1.1, y - 0.02, 2.0, 0.5, who, size=22, bold=True, color=WHITE,
         name=f"next {num} who")
    text(s, MX + 3.1, y, W - 2 * MX - 3.1, 1.1, what, size=18, color=ICE,
         name=f"next {num} what")
footer(s, n, dark=True)

prs.save(OUT)
print(f"wrote {OUT} ({n} slides)")
for w in FIT_WARNINGS:
    print("FIT?", w)
